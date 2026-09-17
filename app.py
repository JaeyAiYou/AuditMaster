import streamlit as st
import pandas as pd
import numpy as np
import io
import pypdf
import google.generativeai as genai
from docx import Document

st.set_page_config(page_title="여신 사후관리 재무분석 & Gemini AI 심사", layout="wide")

# 1. API Key 및 모델 선택 (사이드바)
with st.sidebar:
    st.title("⚙️ AI 및 여신 관리 설정")
    user_api_key = st.text_input("Gemini API Key 입력", type="password", help="Google AI Studio 발급 Key")
    
    model_choice = st.selectbox(
        "Gemini 모델 선택",
        ["gemini-3.7-flash", "gemini-3.6-flash", "gemini-3.1-pro", "직접 입력"],
        index=0
    )
    if model_choice == "직접 입력":
        selected_model_name = st.text_input("사용할 모델명 입력", value="gemini-3.6")
    else:
        selected_model_name = model_choice

st.title("🛡️ 여신 취급 후 사후관리 재무 분석 & AI 사후 진단 시스템")
st.caption("여신 취급 완료 차주의 재무현황 업데이트, 리스크 모니터링 및 향후 사후관리 방안 도출 전용")

# 2. 표준 재무분석표 템플릿 항목 정의 (첨부 양식 기준)
default_items = [
    "자산총계", "유동자산", "현금성자산", "매출채권", "재고자산", "유형자산", "투자자산",
    "부채총계", "유동부채", "매입채무", "차입금",
    "자본총계", "자본금",
    "매출액", "영업이익", "금융비용", "당기순이익", "EBITDA", "영업활동현금흐름",
    "영업이익률(%)", "EBITDA 마진률(%)", "유동비율(%)", "부채비율(%)", "차입금의존도(%)", "EBITDA/금융비용(배)"
]
years_columns = ["2022년", "2023년", "2024년", "2025년", "2026년 반기"]

# 3. 탭 구성
tab_table, tab_files, tab_ai = st.tabs(["📊 표준 재무분석표 (사후관리)", "📁 감사보고서/첨부파일", "🤖 AI 사후관리 방안 & Word 보고서"])

with tab_table:
    st.subheader("1. 여신 사후관리 표준 재무분석표")
    st.info("💡 아래 표에 수치를 직접 수정/입력하거나, [감사보고서/첨부파일] 탭에서 파일 업로드를 병행하세요.")
    
    initial_data = {col: [0.0]*len(default_items) for col in years_columns}
    initial_df = pd.DataFrame(initial_data, index=default_items)
    
    edited_df = st.data_editor(
        initial_df,
        use_container_width=True,
        height=600,
        key="financial_editor"
    )

with tab_files:
    st.subheader("2. 감사보고서 및 주석 첨부 (최대 3개)")
    uploaded_files = st.file_uploader(
        "감사보고서(PDF) 및 재무 자료(XLSX, CSV) 업로드", 
        type=["pdf", "xlsx", "csv"], 
        accept_multiple_files=True
    )
    
    parsed_file_tables = []
    combined_pdf_text = ""
    
    if uploaded_files:
        if len(uploaded_files) > 3:
            st.warning("⚠️ 최대 3개 파일까지 파싱됩니다.")
            uploaded_files = uploaded_files[:3]
            
        for idx, file in enumerate(uploaded_files):
            st.markdown(f"**📄 파일 {idx+1}: {file.name}**")
            if file.name.endswith(".csv"):
                df_up = pd.read_csv(file)
                st.dataframe(df_up, use_container_width=True)
                parsed_file_tables.append((file.name, df_up))
            elif file.name.endswith(".xlsx"):
                df_up = pd.read_excel(file)
                st.dataframe(df_up, use_container_width=True)
                parsed_file_tables.append((file.name, df_up))
            elif file.name.endswith(".pdf"):
                reader = pypdf.PdfReader(file)
                text = "\n".join([p.extract_text() for p in reader.pages if p.extract_text()])
                combined_pdf_text += f"\n--- [{file.name}] ---\n" + text
                st.success(f"PDF 텍스트 추출 완료 ({len(reader.pages)} 페이지)")

with tab_ai:
    st.subheader("3. 여신 사후관리 AI 리스크 진단 및 관리 방안 도출")
    
    if st.button("🚀 사후관리 종합 분석 및 방안 생성"):
        if not user_api_key:
            st.error("🔑 왼쪽 사이드바에 Gemini API Key를 입력해주세요.")
        else:
            try:
                genai.configure(api_key=user_api_key)
                model = genai.GenerativeModel(selected_model_name)
                
                table_str = edited_df.to_string()
                additional_tables_str = ""
                for fname, df_p in parsed_file_tables:
                    additional_tables_str += f"\n[{fname}]\n" + df_p.to_string() + "\n"
                
                prompt = f"""
                당신은 금융기관 수석 여신사후관리역(Risk Manager)입니다.
                본 분석은 **여신 취급이 완료된 차주기업에 대한 사후관리 및 모니터링 목적**입니다.
                제공된 표준 재무분석표와 감사보고서 주석을 바탕으로 여신 사후관리 진단 및 향후 관리 방안을 도출해주세요.

                [1. 표준 재무분석표 데이터]
                {table_str}

                [2. 추가 첨부 파일 재무 데이터]
                {additional_tables_str}

                [3. 감사보고서 주석/텍스트]
                {combined_pdf_text[:12000]}

                [분석 및 보고서 작성 가이드라인]
                여신 사후관리 관점에서 다음 핵심 항목을 포함하여 구체적으로 작성하세요:
                1. **재무현황 변동 추적 및 조기경보 시그널 진단**:
                   - 취급 시점 대비 최근 재무 상태 변화 (매출, 영업이익, 현금흐름, EBITDA/금융비용 등)
                   - 부채비율, 차입금의존도, 유동성 악화 등 조기경보 시그널 점검
                2. **재무약정(Covenant) 및 우발채무 리스크**:
                   - 부채비율 약정 제한선 미준수 여부, 이자보상능력 점검
                   - 주석 상 우발채무, 지급보증, 소송, 담보제공 현황에 따른 위험도
                3. **내부 신용등급 및 관제 단계 평가**:
                   - 여신 관제 단계 (정상 / 조기경보 / 관찰여신 / 부실화 우려 등) 재평가
                4. **향후 구체적 여신 사후관리 방안 (Action Plan)**:
                   - 만기 도래 시 감축 상환 조건 부여 (예: 5~10% 감축상환) 또는 연장 조건
                   - 추가 담보/보증서 요구, 금리 재정산(가산금리), 현장실사 및 영업점 점검 계획
                """
                
                with st.spinner(f"Gemini AI({selected_model_name})가 여신 사후관리 관점에서 진단 중입니다..."):
                    res = model.generate_content(prompt)
                    st.session_state["post_credit_result"] = res.text
                    st.markdown(res.text)
                    
            except Exception as e:
                st.error(f"AI 분석 오류: {e}")

    # Word 보고서 다운로드
    if "post_credit_result" in st.session_state:
        st.divider()
        st.subheader("📑 여신 사후관리 Word 보고서 (.docx)")
        
        doc = Document()
        doc.add_heading('여신 취급 후 사후관리 재무진단 및 관리방안 보고서', level=0)
        
        # 1. 표준 재무분석표 작성
        doc.add_heading('1. 표준 재무분석표 (업데이트 현황)', level=1)
        t = doc.add_table(rows=edited_df.shape[0]+1, cols=edited_df.shape[1]+1)
        t.style = 'Table Grid'
        
        t.rows[0].cells[0].text = "구분"
        for c_idx, col in enumerate(edited_df.columns):
            t.rows[0].cells[c_idx+1].text = str(col)
            
        for r_idx, (idx_label, row) in enumerate(edited_df.iterrows()):
            t.rows[r_idx+1].cells[0].text = str(idx_label)
            for c_idx, val in enumerate(row):
                t.rows[r_idx+1].cells[c_idx+1].text = str(val)
                
        doc.add_paragraph("")
        
        # 2. AI 사후관리 진단 결과
        doc.add_heading('2. 여신 사후관리 AI 리스크 진단 및 향후 관리 방안', level=1)
        doc.add_paragraph(st.session_state["post_credit_result"])
        
        doc_io = io.BytesIO()
        doc.save(doc_io)
        doc_io.seek(0)
        
        st.download_button(
            "📥 여신 사후관리 Word 보고서 다운로드", 
            doc_io, 
            "Post_Credit_Management_Report.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
